import sys
import os
import sqlite3
import pandas as pd
from openpyxl import load_workbook
from datetime import datetime
import warnings

# Silenciar advertencias de formato de openpyxl que ensucian el log
warnings.filterwarnings("ignore", category=UserWarning)

# Importar gestor de configuración
from app_config import get_tax_settings
from utils import setup_console, normalize_category, print_header, print_success, print_error, print_info, get_eastern_now

# Configurar consola (UTF-8)
setup_console()

def process_excel(category, tax_percent=None):
    """
    Actualiza una plantilla de Excel con los precios de la base de datos para una categoría específica.
    Si tax_percent es None, se lee de la configuración de la base de datos.
    """
    try:
        # --- Configuración de Tax ---
        db_tax, db_bypass = get_tax_settings()

        # --- Definición de Rutas ---
        # La ruta base es la carpeta raíz del proyecto 'Catalogo Update Workflow'
        script_dir = os.path.dirname(os.path.abspath(__file__))
        project_root = os.path.dirname(script_dir) # Esto nos lleva a la carpeta principal del proyecto
        db_path = os.path.join(project_root, 'database', 'diccionario.db') # Ruta correcta a la DB
        template_path = os.path.join(project_root, 'data', 'excel_templates', f'Plantilla_{category}.xlsx')
        output_dir = os.path.join(project_root, 'data', 'excel_output')

        # Crea el directorio de salida si no existe
        os.makedirs(output_dir, exist_ok=True)

        # Genera un nombre de archivo único con timestamp en hora de Orlando Florida
        timestamp = get_eastern_now().strftime("%Y%m%d_%H%M%S")
        output_filename = f'Plantilla_{category}_Actualizada_{timestamp}.xlsx'
        output_path = os.path.join(output_dir, output_filename)

        # --- Verificación de Archivos ---
        if not os.path.exists(db_path):
            print_error(f"[ERROR] La base de datos no se encontró en: {db_path}")
            return
        if not os.path.exists(template_path):
            print_error(f"[ERROR] La plantilla de Excel no se encontró en: {template_path}")
            return

        # --- Carga de Precios desde SQLite ---
        print_header("GENERANDO EXCEL FINAL")
        print_info("Leyendo precios actualizados de la base de datos...")
        with sqlite3.connect(db_path, timeout=30.0) as conn:
            query = """
                SELECT
                    mp.nombre_procore,
                    cw.precio_actual
                FROM mapeo_procore mp
                JOIN catalogo_web cw ON mp.web_item_id = cw.id
                WHERE cw.precio_actual > 0 AND cw.categoria = ?;
            """
            # Usar pandas para leer la consulta y crear un diccionario para búsqueda rápida
            df = pd.read_sql_query(query, conn, params=(category,))
            price_map = pd.Series(df.precio_actual.values, index=df.nombre_procore).to_dict()

        if not price_map:
            print_error(f"[ADVERTENCIA] No hay precios para '{category}' en la BD. El Excel saldrá vacío.")
        else:
            print_success(f"✓ {len(price_map)} precios cargados en memoria.")

        # --- Lógica de Tax Efectiva ---
        if db_bypass:
            effective_tax = 0.0
            print_info("MODO BYPASS ACTIVO: Se aplicará 0% de impuesto.")
        else:
            # Si se pasó un argumento explícito (no None), úsalo. Si no, usa el de la BD.
            if tax_percent is not None:
                effective_tax = tax_percent
            else:
                effective_tax = db_tax

        # --- Procesamiento del Archivo Excel ---
        print_info(f"Procesando plantilla: {os.path.basename(template_path)}")
        wb = load_workbook(template_path)
        ws = wb.active

        # --- Configuración Fija de Columnas ---
        # Según estructura de plantilla Procore:
        desc_col_idx = 4 # Columna D (Description / Name)
        cost_col_idx = 7 # Columna G (Unit Cost)
        header_row = 1   # Los datos comienzan en la fila 2
        
        print_success(f"✓ Usando mapeo fijo: Col D ({desc_col_idx}) para Nombre y Col G ({cost_col_idx}) para Costo")

        updated_count = 0
        # Itera sobre todas las filas a partir de la segunda (para saltar el encabezado)
        for row in ws.iter_rows(min_row=header_row + 1, values_only=False):
            procore_name_cell = row[desc_col_idx - 1]
            procore_name = procore_name_cell.value

            # Si el nombre de Procore existe en nuestro mapa de precios...
            if procore_name and procore_name in price_map:
                # Obtenemos el precio original de la base de datos
                original_price = price_map[procore_name]
                
                # Calculamos el multiplicador del impuesto (ej. 6.5% -> 1.065)
                tax_multiplier = 1 + (effective_tax / 100.0)
                
                # Aplicamos el impuesto y redondeamos a 2 decimales para formato de moneda
                final_price = round(original_price * tax_multiplier, 2)
                # Actualizamos la celda de 'Unit Cost' con el precio final (con impuesto).
                # Esta operación solo cambia el valor, manteniendo el formato existente de la celda.
                ws.cell(row=procore_name_cell.row, column=cost_col_idx, value=final_price)
                updated_count += 1

        print_success(f"✓ Se actualizaron {updated_count} filas en el Excel.")
        print_info(f"Tax aplicado: {effective_tax}%")

        # --- Guardado del Archivo ---
        wb.save(output_path)
        # Usamos [EXITO] sin caracteres especiales para evitar problemas de codificación en Windows
        print(f"\n<div class='log-success'>[EXITO] Archivo actualizado guardado en: {output_path}</div>")

    except PermissionError as e:
        print_error("¡Error de Acceso! El archivo está bloqueado.")
        print_info(f"Archivo bloqueado: {e.filename}")
        print_info("Sugerencia: Cierra el Excel y asegúrate de que OneDrive no esté sincronizando el archivo.")
    except Exception as e:
        print_error(f"[ERROR FATAL] {e}")

if __name__ == "__main__":
    if len(sys.argv) < 2:
        print(f"Uso: python procesar_excel.py <categoria>")
        print(f"Uso opcional: python procesar_excel.py <categoria> <tax_percent>")
        sys.exit(1)

    category_arg = sys.argv[1]
    
    # Normalizar para consistencia con la DB centralizando la lógica
    category_arg = normalize_category(category_arg)

    # Capturar el porcentaje de tax si viene como argumento
    # Si no se pasa argumento, se envía None para que la función lea la config global
    tax_arg = None 
    if len(sys.argv) > 2:
        try:
            tax_arg = float(sys.argv[2])
        except ValueError:
            print(f"[ADVERTENCIA] Tax '{sys.argv[2]}' inválido. Usando configuración guardada.")

    process_excel(category_arg, tax_arg)